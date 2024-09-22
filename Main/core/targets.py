from io import BytesIO
from re import sub

from django.core.exceptions import ObjectDoesNotExist
from django.core.validators import validate_email
from django.db import transaction
from django.forms import ValidationError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment

from Main.models import Target, TargetList, Attribute


def load_targets_from_excel(targets_list: TargetList, uploaded_file):
    wb = load_workbook(filename=uploaded_file, read_only=True)
    ws = wb.active
    targets = {}
    header = next(ws.rows)
    tags_keys = []
    n = 1
    for x in header[1:]:
        if x.value is not None:
            tags_keys.append('ORDN-' + "%03d" % n + '-' + x.value)
            n += 1
    for row in ws.iter_rows(min_row=2):
        if row[0] and row[0].value:
            email = row[0].value.lower()
            tags_values = [x.value if x.value is not None else "N/A" for x in row[1:]]
            tags = dict(zip(tags_keys, tags_values))
            try:
                validate_email(email)
            except ValidationError:
                raise ValidationError(email)
            targets[email] = tags
    new_targets = []
    with transaction.atomic():
        for email in targets:
            if not targets_list.targets.filter(mail_address=email).exists():
                tmp = Target.objects.create(mail_address=email)
                atts = []
                for tag, value in targets[email].items():
                    try:
                        new_att = Attribute.objects.get(key=tag, value=value)
                    except ObjectDoesNotExist:
                        new_att = Attribute.objects.create(key=tag, value=value)
                    atts.append(new_att)
                tmp.attributes.add(*atts)
                new_targets.append(tmp)
        targets_list.targets.add(*new_targets)

    transaction.commit()

    return True


def export_to_xlsx(targets_list: TargetList):
    targets = targets_list.targets.all()
    wb = Workbook()
    ws = wb.active
    tags = set()
    mails = {}
    for target in targets:
        mails[target.mail_address] = {}
        attributes = target.attributes.all()
        for attribute in attributes:
            mails[target.mail_address][attribute.key] = attribute.value
            tags.add(attribute.key)
    raw_header = list(sorted(tags))
    header = []
    for h in raw_header:
        h = sub(r'ORDN-[0-9]{3}-', '', h)
        header.append(h)
    header.insert(0, "email")
    ws.append(header)
    for address in mails:
        line = [address]
        for tag in sorted(tags):
            if tag in mails[address]:
                line.append(mails[address][tag])
            else:
                line.append("N/A")
        ws.append(line)

    c = ws['B2']
    ws.freeze_panes = c
    head = ws[1]
    ft = Font(bold=True)
    al = Alignment(horizontal="center", vertical="center")
    for cell in head:
        cell.font = ft
    for row in ws.rows:
        for cell in row:
            cell.alignment = al
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
