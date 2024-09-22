from datetime import datetime
from logging import getLogger
from re import sub
from smtplib import SMTPRecipientsRefused, SMTPServerDisconnected
from time import sleep

from django.core import mail
from django.utils.timezone import get_current_timezone
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from Main.core.mail import build_email, build_attachment, send_email, send_email_with_attachment
from Main.models import Campaign, AnonymousTarget
from Main.utils import send_alert_new_campaign


def start(campaign: Campaign):
    logger = getLogger(__name__)
    if campaign.status != "2":
        campaign.status = "2"
    attachment_content = build_attachment(campaign)
    campaign.targets_count = campaign.count_targets()
    campaign.save()
    try:
        send_alert_new_campaign(campaign.targets_count,
                                campaign.author.user.email,
                                campaign.mail_template.title,
                                "%s@%s" % (campaign.from_name, campaign.from_domain.domain)
                                )
    except SMTPRecipientsRefused:
        pass
    targetlists = campaign.targets.all()
    con = mail.get_connection(fail_silently=False, timeout=15)
    con.open()
    mail_content = build_email(campaign)
    for targetlist in targetlists:
        targets = targetlist.targets.all()
        for target in targets:
            new_anon = AnonymousTarget()
            new_anon.save()
            new_anon.import_attributes(target)
            try:
                if campaign.campaign_type in ["1", "3", "4"]:
                    send_email(campaign, target.mail_address, new_anon.uniqueid, con, mail_content)
                elif campaign.campaign_type == "2":
                    send_email_with_attachment(campaign, target.mail_address, new_anon.uniqueid, con, mail_content,
                                               attachment_content)
            except SMTPRecipientsRefused:
                pass
            except SMTPServerDisconnected:
                logger.error("Timeout connecting to SMTP server")
                logger.info("Pause for a little while before opening a new connection")
                sleep(60)
                try:
                    con = mail.get_connection(fail_silently=False, timeout=15)
                    con.open()
                    if campaign.campaign_type in ["1", "3", "4"]:
                        send_email(campaign, target.mail_address,
                                   new_anon.uniqueid,
                                   con,
                                   mail_content
                                   )
                    elif campaign.campaign_type == "2":
                        send_email_with_attachment(campaign, target.mail_address,
                                                   new_anon.uniqueid,
                                                   con,
                                                   mail_content,
                                                   attachment_content
                                                   )
                except SMTPServerDisconnected:
                    logger.error("Failed to reopen a connection to SMTP server. Giving up")
                    break
                else:
                    logger.info("New connection is working")
                    new_anon.mail_sent_time = datetime.now(tz=get_current_timezone())
                    campaign.anonymous_targets.add(new_anon)
                    new_anon.save()
            else:
                new_anon.mail_sent_time = datetime.now(tz=get_current_timezone())
                campaign.anonymous_targets.add(new_anon)
                new_anon.save()
    con.close()
    return True


def stop(campaign: Campaign):
    campaign.status = "3"
    campaign.save()
    return True


def test(campaign: Campaign, recipient):
    con = mail.get_connection()
    con.open()
    mail_content = build_email(campaign)
    attachment_content = build_attachment(campaign)
    if campaign.campaign_type == "1" or campaign.campaign_type == "3" or campaign.campaign_type == "4":
        send_email(campaign, recipient, campaign.testid, con, mail_content)
    elif campaign.campaign_type == "2":
        send_email_with_attachment(campaign, recipient, campaign.testid, con, mail_content, attachment_content)
    con.close()
    return True


def generate_results_xlsx(campaign: Campaign):
    targets = campaign.anonymous_targets.all()
    wb = Workbook()
    ws = wb.active
    tags = set()
    dest_filename = 'results/' + str(campaign.id) + '.xlsx'

    for target in targets:
        attributes = target.attributes.all()
        for attribute in attributes:
            tags.add(attribute.key)

    header = ["id", "mail sent time", "mail opened", "mail opened time"]
    links_autoclicked = campaign.links_autoclicked()

    if campaign.campaign_type == "1" or campaign.campaign_type == "3" or campaign.campaign_type == "4":
        header.append("link clicked")
        header.append("link clicked time")
        if links_autoclicked:
            header.append("link autoclicked time")
        if campaign.campaign_type == "3":
            header.append("form submitted")
            header.append("form submitted time")
        if campaign.campaign_type == "4":
            header.append("img clicked")
            header.append("img clicked time")
    else:
        header.append("attachment opened")
        header.append("attachment opened time")
    header.append("reported")
    header.append("reported time")

    for tag in sorted(tags):
        header.append(sub(r'ORDN-[0-9]{3}-', '', tag))

    ft = Font(bold=True)
    al = Alignment(horizontal="center", vertical="center")
    al2 = Alignment(horizontal="left", vertical="center")
    column = 1
    for cell in header:
        ws.cell(row=1, column=column, value=cell)
        ws.cell(row=1, column=column).font = ft
        ws.cell(row=1, column=column).alignment = al
        column += 1

    row = 2
    for target in targets:
        values = []

        if target.mail_sent_time:
            mail_sent_time = target.mail_sent_time.strftime('%Y-%m-%d %H:%M:%S')
        else:
            mail_sent_time = "N/A"

        if target.mail_opened:
            mail_opened = "yes"
        else:
            mail_opened = "no"

        if target.mail_opened_time:
            mail_opened_time = target.mail_opened_time.strftime('%Y-%m-%d %H:%M:%S')
        else:
            mail_opened_time = "N/A"

        if target.link_clicked:
            link_clicked = "yes"
        else:
            link_clicked = "no"

        if target.link_clicked_time:
            link_clicked_time = target.link_clicked_time.strftime('%Y-%m-%d %H:%M:%S')
        else:
            link_clicked_time = "N/A"
        if target.autoclick_time:
            link_autoclick_time = target.autoclick_time.strftime('%Y-%m-%d %H:%M:%S')
        else:
            link_autoclick_time = "N/A"

        if target.attachment_opened:
            attachment_opened = "yes"
        else:
            attachment_opened = "no"

        if target.attachment_opened_time:
            attachment_opened_time = target.attachment_opened_time.strftime('%Y-%m-%d %H:%M:%S')
        else:
            attachment_opened_time = "N/A"

        if target.form_submitted:
            form_submitted = "yes"
        else:
            form_submitted = "no"

        if target.form_submitted_time:
            form_submitted_time = target.form_submitted_time.strftime('%Y-%m-%d %H:%M:%S')
        else:
            form_submitted_time = "N/A"

        if target.reported:
            reported = "yes"
        else:
            reported = "no"

        if target.reported_time:
            reported_time = target.reported_time.strftime('%Y-%m-%d %H:%M:%S')
        else:
            reported_time = "N/A"

        for tag in sorted(tags):
            att = target.attributes.filter(key=tag)
            if att:
                values.append(att[0].value)
            else:
                values.append("N/A")

        column = 1
        ws.cell(row=row, column=column, value=target.uniqueid)
        ws.cell(row=row, column=column).font = ft
        ws.cell(row=row, column=column).alignment = al2
        column += 1
        ws.cell(row=row, column=column, value=mail_sent_time)
        ws.cell(row=row, column=column).alignment = al
        column += 1
        ws.cell(row=row, column=column, value=mail_opened)
        ws.cell(row=row, column=column).alignment = al
        column += 1
        ws.cell(row=row, column=column, value=mail_opened_time)
        ws.cell(row=row, column=column).alignment = al
        column += 1
        if campaign.campaign_type == "1" or campaign.campaign_type == "3" or campaign.campaign_type == "4":
            ws.cell(row=row, column=column, value=link_clicked)
            ws.cell(row=row, column=column).alignment = al
            column += 1
            ws.cell(row=row, column=column, value=link_clicked_time)
            ws.cell(row=row, column=column).alignment = al
            column += 1
            if links_autoclicked:
                ws.cell(row=row, column=column, value=link_autoclick_time)
                ws.cell(row=row, column=column).alignment = al
                column += 1
            if campaign.campaign_type == "3" or campaign.campaign_type == "4":
                ws.cell(row=row, column=column, value=form_submitted)
                ws.cell(row=row, column=column).alignment = al
                column += 1
                ws.cell(row=row, column=column, value=form_submitted_time)
                ws.cell(row=row, column=column).alignment = al
                column += 1
        else:
            ws.cell(row=row, column=column, value=attachment_opened)
            ws.cell(row=row, column=column).alignment = al
            column += 1
            ws.cell(row=row, column=column, value=attachment_opened_time)
            ws.cell(row=row, column=column).alignment = al
            column += 1
        ws.cell(row=row, column=column, value=reported)
        ws.cell(row=row, column=column).alignment = al
        column += 1
        ws.cell(row=row, column=column, value=reported_time)
        ws.cell(row=row, column=column).alignment = al
        column += 1
        for val in values:
            ws.cell(row=row, column=column, value=val)
            ws.cell(row=row, column=column).alignment = al
            column += 1
        row += 1

    c = ws['B2']
    ws.freeze_panes = c
    wb.save(filename=dest_filename)
    return True
